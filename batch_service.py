"""
Batch company research.

Companies are processed ONE AT A TIME through a controlled queue. Each company's
result and PDF are written the moment that company finishes, so reports are
downloadable while the rest of the list is still running, and an interrupted batch
resumes from where it stopped.

Within a single company, "Run All Three" keeps the existing behaviour: retrieval
runs once and the models synthesise in parallel over that shared evidence.
"""

import concurrent.futures
import csv
import io
import json
import re
import shutil
import threading
import time
from datetime import datetime
from pathlib import Path

import confidence as conf
import language_view as lv
import pdf_service as ps
import research_service as rs

HERE = Path(__file__).resolve().parent
REPORTS_DIR = HERE / "reports"
BATCH_DIR = REPORTS_DIR / "_batches"

NAME_HINTS = ("company name", "company", "name", "公司名称", "公司", "客户", "企业名称")
SITE_HINTS = ("company website", "website", "url", "site", "domain", "官网", "网站", "网址")

STATUS_PENDING = "Pending"
STATUS_SEARCHING = "Searching"
STATUS_GENERATING = "Generating"
STATUS_PDF = "PDF Generating"
STATUS_COMPLETE = "Completed"
STATUS_FAILED = "Failed"
STATUS_TIMEOUT = "Timed Out"
# Bumped when the retrieval pipeline or report structure changes materially, so
# a saved report can be told apart from one the current engine would produce.
#   1 = single-query retrieval, 12-section report
#   2 = multi-area retrieval + official-site discovery + 19-section SKEQI report
CURRENT_REPORT_VERSION = 2
CURRENT_MIN_QUERIES = 15
CURRENT_MIN_SECTIONS = 19
MIN_HEALTHY_SOURCES = 5

STATUS_SKIPPED = "Existing Report"   # 已有报告 — a saved report was reused, no AI call

ALL_STATUSES = (STATUS_PENDING, STATUS_SEARCHING, STATUS_GENERATING, STATUS_PDF,
                STATUS_COMPLETE, STATUS_FAILED, STATUS_TIMEOUT, STATUS_SKIPPED)


# --------------------------------------------------------------- file parsing
def parse_upload(filename, data):
    """Return (headers, rows) from a .csv or .xlsx upload."""
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = [[("" if c is None else str(c).strip()) for c in r]
                for r in ws.iter_rows(values_only=True)]
        wb.close()
    else:
        text = None
        for enc in ("utf-8-sig", "utf-8", "gb18030", "latin-1"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text or ""))]
    rows = [r for r in rows if any(c for c in r)]
    if not rows:
        return [], []
    return rows[0], rows[1:]


_URLISH = re.compile(r"^(https?://|www\.)|\.(com|cn|net|org|io|co|de|jp|kr|eu|us|uk)\b", re.I)


def _looks_like_websites(rows, idx, sample=12):
    """True when a column's values are mostly URLs or bare domains."""
    if idx is None:
        return False
    seen = hits = 0
    for row in rows[:sample]:
        v = str((row[idx] if idx < len(row) else "") or "").strip()
        if not v:
            continue
        seen += 1
        if _URLISH.search(v):
            hits += 1
    return seen > 0 and hits >= max(1, int(seen * 0.6))


def guess_mapping(headers, rows=None):
    """Detect the company and website columns, and say how sure we are.

    The caller needs the confidence, not just the guess: asking every user to
    confirm two dropdowns on every upload is noise when the header literally
    says "Company Name", and silently guessing column 0 is wrong when it does
    not. Only a genuinely ambiguous company column should prompt anyone.
    """
    rows = rows or []
    lower = [str(h or "").lower().strip() for h in headers]
    name_idx = site_idx = None
    for i, h in enumerate(lower):
        if name_idx is None and any(k == h or k in h for k in NAME_HINTS):
            name_idx = i
        if site_idx is None and any(k == h or k in h for k in SITE_HINTS):
            site_idx = i

    name_confident = name_idx is not None
    site_confident = site_idx is not None
    # An unlabelled column that is full of URLs is the website column anyway.
    if site_idx is None:
        for i in range(len(headers)):
            if i != name_idx and _looks_like_websites(rows, i):
                site_idx, site_confident = i, True
                break
    # Never fall back to "column 1 is probably the website". A missing website
    # is fine — research resolves the official site itself — and a wrong one
    # sends the whole batch to the wrong domains.
    if name_idx is None:
        name_idx = 0

    return {"name": name_idx, "website": site_idx,
            "name_confident": name_confident, "website_confident": site_confident,
            "ambiguous": not name_confident}


ITEM_DEFAULTS = {
    "status": STATUS_PENDING, "model": None, "sources": None, "tokens": 0,
    "input_tokens": 0, "output_tokens": 0, "latency": None, "seconds": None,
    "error": None, "error_reason": None, "note": None,
    "pdf": None, "pdfs": [], "cost": None, "selected": True,
    "website_status": None, "resolved_website": None, "ignore_website": False,
    "report_version": None, "limited_sources": False,
    "trust_website": False, "force": False, "evidence_level": None,
    "report_dir": None,
}


def normalize_item(raw):
    """Coerce one client-supplied batch item into the full internal shape.

    The payload is never trusted just because the current UI happens to produce
    the expected structure: missing keys get safe defaults and wrong types are
    corrected, so downstream code cannot raise KeyError/TypeError.
    Returns None when the item carries no usable company name.
    """
    if not isinstance(raw, dict):
        return None
    company = str(raw.get("company") or raw.get("name") or "").strip()
    if not company:
        return None
    website = str(raw.get("website") or raw.get("url") or "").strip()
    if website and not re.match(r"^https?://", website, re.I):
        website = "https://" + website.lstrip("/")

    item = dict(ITEM_DEFAULTS)
    item["pdfs"] = []
    item["company"] = company[:200]
    item["website"] = website[:500]

    item["ignore_website"] = bool(raw.get("ignore_website"))
    item["trust_website"] = bool(raw.get("trust_website"))
    item["force"] = bool(raw.get("force"))       # explicit "Generate Anyway"
    for key in ("status", "model", "error", "pdf", "error_reason", "note",
                "website_status", "resolved_website", "report_dir"):
        val = raw.get(key)
        item[key] = str(val)[:300] if isinstance(val, (str, int, float)) and val != "" else ITEM_DEFAULTS[key]
    if item["status"] not in ALL_STATUSES:
        item["status"] = STATUS_PENDING

    item["limited_sources"] = bool(raw.get("limited_sources"))
    for key in ("tokens", "input_tokens", "output_tokens", "sources", "report_version"):
        val = raw.get(key)
        try:
            item[key] = int(val) if val is not None else ITEM_DEFAULTS[key]
        except (TypeError, ValueError):
            item[key] = ITEM_DEFAULTS[key]

    for key in ("latency", "seconds", "cost"):
        val = raw.get(key)
        try:
            item[key] = float(val) if val is not None else None
        except (TypeError, ValueError):
            item[key] = None

    pdfs = raw.get("pdfs")
    if isinstance(pdfs, list):
        for entry in pdfs:
            if isinstance(entry, dict) and entry.get("file"):
                item["pdfs"].append({"model": str(entry.get("model") or "")[:80],
                                     "file": str(entry["file"])[:200]})

    item["selected"] = bool(raw.get("selected", True))
    return item


def normalize_items(raw_items):
    """Returns (items, skipped_count). Malformed entries are dropped, not fatal."""
    if not isinstance(raw_items, list):
        return [], 0
    items, skipped, seen = [], 0, set()
    for raw in raw_items:
        item = normalize_item(raw)
        if item is None:
            skipped += 1
            continue
        key = item["company"].lower()
        if key in seen:
            skipped += 1
            continue
        seen.add(key)
        items.append(item)
    return items, skipped


def build_items(rows, mapping):
    items, seen = [], set()
    ni, si = mapping.get("name"), mapping.get("website")
    for r in rows:
        name = (r[ni].strip() if ni is not None and ni < len(r) else "")
        site = (r[si].strip() if si is not None and si < len(r) else "")
        if not name:
            continue
        if site and not re.match(r"^https?://", site, re.I):
            site = "https://" + site.lstrip("/")
        key = name.lower()
        if key in seen:
            continue
        seen.add(key)
        items.append({"company": name, "website": site, "status": STATUS_PENDING,
                      "model": None, "sources": None, "tokens": None, "seconds": None,
                      "pdfs": [], "error": None, "selected": True})
    return items


# --------------------------------------------------------------- report store
def company_dir(company):
    d = REPORTS_DIR / ps.safe_name(company)
    d.mkdir(parents=True, exist_ok=True)
    return d


def existing_report(company, model=None):
    """Completed report for this company (and model, if given), else None."""
    d = REPORTS_DIR / ps.safe_name(company)
    meta = d / "_meta.json"
    if not meta.exists():
        return None
    try:
        data = json.loads(meta.read_text(encoding="utf-8"))
    except Exception:
        return None
    runs = data.get("runs") or []
    if model:
        runs = [r for r in runs if r.get("model") == model]
    if not runs:
        return None
    return data


def report_version(company_dir_path, run):
    """Which pipeline produced this saved report?

    New reports carry the version explicitly. Older ones predate the marker, so
    it is inferred from what the file actually contains: the 19-section report
    and the multi-area query plan only exist in version 2.
    """
    if run.get("report_version"):
        return int(run["report_version"])
    rec_path = Path(company_dir_path) / run.get("json", "research.json")
    if not rec_path.exists():
        rec_path = Path(company_dir_path) / "research.json"
    try:
        rec = json.loads(rec_path.read_text(encoding="utf-8"))
    except Exception:
        return 1
    sections = len(re.findall(r"^##\s", rec.get("research_result") or "", re.M))
    queries = len(rec.get("search_queries") or [])
    return 2 if (sections >= CURRENT_MIN_SECTIONS and queries >= CURRENT_MIN_QUERIES) else 1


def _archive_previous(d, run_model):
    """Keep the superseded report rather than deleting it on regenerate."""
    prior = d / "research_{}.json".format(run_model)
    if not prior.exists():
        return
    hist = d / "history"
    hist.mkdir(exist_ok=True)
    try:
        ts = datetime.fromtimestamp(prior.stat().st_mtime).strftime("%Y%m%d-%H%M%S")
        shutil.copy2(prior, hist / "research_{}_{}.json".format(run_model, ts))
        meta_p = d / "_meta.json"
        if meta_p.exists():
            meta = json.loads(meta_p.read_text(encoding="utf-8"))
            for r in meta.get("runs", []):
                old_pdf = d / r.get("pdf", "")
                if r.get("model") == run_model and old_pdf.exists():
                    shutil.copy2(old_pdf, hist / "{}_{}".format(ts, old_pdf.name))
    except Exception:
        pass          # history is a convenience; never block a new report on it


def write_company_output(company, website, package, run):
    """Write research JSON + bilingual PDF for one completed model run."""
    d = company_dir(company)
    _archive_previous(d, run["model"])
    stamp = datetime.now().astimezone()
    record = {
        "timestamp": stamp.isoformat(),
        "company": company, "website": website,
        "model": run["model"], "model_label": run["model_label"],
        "protocol": run["protocol"], "endpoint": run["endpoint"],
        "search_enabled": True,
        "evidence_cached": package.get("cached", False),
        "search_queries": package["search_queries"],
        "search_requests": len(package["search_queries"]) if not package.get("cached") else 0,
        "sources": package["sources"],
        "research_result": conf.normalize_text(run["report"]),
        "research_result_display": conf.strip_unsupported(run["report"]),
        "research_result_raw": run["report"],
        "confidence": conf.extract(run["report"]),
        "token_usage": {"input": run["input_tokens"], "output": run["output_tokens"],
                        "total": run["total_tokens"]},
        "latency_seconds": run["latency_seconds"],
        "retrieval_timings": package.get("timings", {}),
        "report_version": CURRENT_REPORT_VERSION,
        "limited_evidence": bool(package.get("limited_evidence")),
        "financial_sources": package.get("financial_sources", {}),
        "apollo_usage": (package.get("apollo") or {}).get("usage", {}),
        "quality": package.get("quality", {}),
        "decision_makers": run.get("decision_makers", []),
        "people_summary": run.get("people_summary", {}),
        "status": run["status"], "error": run["error"],
    }
    multi = True
    json_name = "research_%s.json" % run["model"]
    (d / json_name).write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding="utf-8")
    # Canonical single-file name for convenience.
    (d / "research.json").write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding="utf-8")

    # The canonical export stays bilingual; EN/ZH are rendered on demand from the
    # same stored record, so switching language never re-runs research.
    pdf_name = ps.pdf_filename(company, run["model_label"], stamp, multi_model=multi,
                               lang=lv.BILINGUAL)
    ps.build_pdf(record, d / pdf_name, lang=lv.BILINGUAL)

    meta_path = d / "_meta.json"
    meta = {"company": company, "website": website, "runs": []}
    if meta_path.exists():
        try:
            meta = json.loads(meta_path.read_text(encoding="utf-8"))
        except Exception:
            pass
    meta["company"], meta["website"] = company, website
    meta["updated_at"] = stamp.isoformat()
    meta["runs"] = [r for r in meta.get("runs", []) if r.get("model") != run["model"]]
    meta["runs"].append({
        "model": run["model"], "model_label": run["model_label"],
        "json": json_name, "pdf": pdf_name,
        "sources": len(package["sources"]),
        "report_version": CURRENT_REPORT_VERSION,
        "limited_evidence": bool(package.get("limited_evidence")),
        "token_usage": record["token_usage"],
        "latency_seconds": run["latency_seconds"],
        "search_requests": record["search_requests"],
        "timestamp": stamp.isoformat(),
    })
    meta_path.write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
    return record, pdf_name


# --------------------------------------------------------------- pricing
def pricing(cfg):
    """RETIRED. Always returns None; the CRM is the single pricing source.

    This was a second, divergent price card living in engine env vars
    (AI_PRICE_INPUT_PER_1K / AI_PRICE_OUTPUT_PER_1K). The CRM already holds a
    versioned, per-(provider, model) pricing table with effective dates, and all
    Account Research cost is now calculated there from the provider's own usage
    blocks. Two cards can disagree, and the one nobody maintains is the one that
    will be wrong.

    It is neutered rather than deleted because its three call sites are in an
    ACTIVE workflow (batch runs and the portfolio PDF) and all of them already
    handle None by omitting cost. In production AI_PRICE_* was never set, so this
    has always returned None and no displayed number changes.
    """
    return None


def estimate_cost(usage, price):
    if not price or not usage:
        return None
    return round((usage.get("input") or 0) / 1000.0 * price["input_per_1k"]
                 + (usage.get("output") or 0) / 1000.0 * price["output_per_1k"], 4)


# --------------------------------------------------------------- batch runner
class Batch:
    """One batch run. Sequential over companies; parallel across models within one."""

    def __init__(self, batch_id, items, models, use_existing=True):
        self.id = batch_id
        self.items = items
        self.models = models
        self.use_existing = use_existing
        self.status = "queued"          # queued | running | stopped | done
        self.current = None
        self.stage = ""
        self.started_at = time.time()
        self.elapsed_completed = 0.0
        self.stop_requested = False
        self.lock = threading.Lock()

    # ---- persistence (so an interrupted batch can resume) ----
    def path(self):
        BATCH_DIR.mkdir(parents=True, exist_ok=True)
        return BATCH_DIR / ("%s.json" % self.id)

    def save(self):
        self.path().write_text(json.dumps(self.snapshot(persist=True),
                                          ensure_ascii=False, indent=2), encoding="utf-8")

    def snapshot(self, persist=False):
        done = [i for i in self.items if i["status"] == STATUS_COMPLETE]
        failed = [i for i in self.items if i["status"] in (STATUS_FAILED, STATUS_TIMEOUT)]
        pending = [i for i in self.items
                   if i["status"] not in (STATUS_COMPLETE, STATUS_FAILED,
                                          STATUS_TIMEOUT, STATUS_SKIPPED)]
        secs = [i["seconds"] for i in done if i.get("seconds")]
        toks = [i["tokens"] for i in done if i.get("tokens")]
        avg_s = round(sum(secs) / len(secs), 1) if secs else None
        elapsed = round(time.time() - self.started_at, 1) if self.status == "running" else None
        snap = {
            "id": self.id, "status": self.status, "current": self.current, "stage": self.stage,
            "models": self.models, "items": self.items,
            "completed": len(done), "failed": len(failed),
            "remaining": len(pending), "total": len(self.items),
            "total_tokens": sum(toks) if toks else 0,
            "avg_tokens": int(sum(toks) / len(toks)) if toks else None,
            "avg_seconds": avg_s,
            "elapsed_seconds": elapsed,
            "eta_seconds": int(avg_s * len(pending)) if (avg_s and pending) else None,
        }
        if persist:
            snap["saved_at"] = datetime.now().astimezone().isoformat()
        return snap

    # ---- execution ----
    def run(self, cfg):
        self.status = "running"
        self.save()
        price = pricing(cfg)
        for item in self.items:
            if self.stop_requested:
                self.status = "stopped"
                self.save()
                return
            if not item.get("selected"):
                continue
            if item["status"] in (STATUS_COMPLETE, STATUS_SKIPPED):
                continue

            self.current = item["company"]
            started = time.time()
            try:
                self._run_one(item, cfg, price)
            except rs.RetrievalError as e:               # a named, actionable reason
                item["status"] = STATUS_FAILED
                item["error"] = str(e)[:300]
                item["error_reason"] = e.reason
                item["website_status"] = getattr(e, "website_status", "unverified")
                item["resolved_website"] = getattr(e, "website", "") or ""
            except Exception as e:                       # one company must not stop the batch
                item["status"] = STATUS_FAILED
                item["error"] = str(e)[:300]
                item["error_reason"] = "timeout" if "timeout" in str(e).lower() else None
            item["seconds"] = round(time.time() - started, 1)
            self.save()

        self.current = None
        self.stage = ""
        self.status = "done" if not self.stop_requested else "stopped"
        self.save()

    def _run_one(self, item, cfg, price):
        company, website = item["company"], item["website"]

        # 12. Avoid duplicate research unless a refresh was requested.
        if self.use_existing:
            prior = existing_report(company)
            wanted = set(self.models)
            have = {r["model"] for r in (prior or {}).get("runs", [])}
            if prior and wanted.issubset(have):
                runs = [r for r in prior["runs"] if r["model"] in wanted]
                versions = [report_version(REPORTS_DIR / ps.safe_name(company), r) for r in runs]
                ver = min(versions) if versions else 1
                src_count = runs[0].get("sources") or 0
                item.update(status=STATUS_SKIPPED,
                            report_version=ver,
                            limited_sources=src_count < MIN_HEALTHY_SOURCES,
                            note="Existing report found — no AI call made. / 已有报告，未调用模型。",
                            model=", ".join(r["model_label"] for r in runs),
                            sources=runs[0].get("sources"),
                            tokens=sum((r.get("token_usage") or {}).get("total") or 0 for r in runs),
                            pdfs=[{"model": r["model_label"], "file": r["pdf"]} for r in runs],
                            report_dir=ps.safe_name(company),
                            resolved_website=prior.get("website") or website,
                            website_status="provided" if (prior.get("website") or website) else None,
                            error=None, error_reason=None)
                return

        item["status"] = STATUS_SEARCHING
        item["note"] = None
        self.stage = "Searching web..."
        self.save()

        hint = "" if item.get("ignore_website") else website
        package = rs.load_cached_evidence(company, hint) if self.use_existing else None
        if not package:
            package = rs.build_shared_evidence(
                company, hint, cfg,
                progress=lambda _s, msg, **_k: setattr(self, "stage", msg),
                trust_website=bool(item.get("trust_website")) and not item.get("ignore_website"))

        item["website_status"] = package.get("website_status")
        item["resolved_website"] = package.get("website") or ""
        if item["website_status"] == "auto_discovered" and item["resolved_website"]:
            item["website"] = item["resolved_website"]

        item["sources"] = len(package["sources"])
        quality = package.get("quality", {})
        item["evidence_level"] = quality.get("level")
        # Same guard as single-company mode: a company with no company-specific
        # source is flagged for review instead of consuming tokens on a report
        # that could only say "not enough evidence".
        if quality.get("blocking") and not item.get("force"):
            item.update(status=STATUS_FAILED,
                        error="Research retrieval incomplete - "
                              + "; ".join(quality.get("reasons", [])[:2]),
                        error_reason="retrieval_incomplete")
            return
        item["status"] = STATUS_GENERATING
        self.stage = "Generating research..."
        self.save()

        # Shared evidence, parallel synthesis - unchanged from single-company mode.
        def one(model):
            try:
                return rs.synthesize_with_fallback(
                    model, company, website, package["evidence"], cfg,
                    apollo_people=(package.get("apollo") or {}).get("people"),
                    progress=lambda m: setattr(self, "stage", m))
            except Exception as e:
                return {"model": model, "model_label": rs.MODEL_LABELS.get(model, model),
                        "status": 0, "report": "", "endpoint": "?",
                        "protocol": "DashScope Native", "input_tokens": None,
                        "output_tokens": None, "total_tokens": None,
                        "latency_seconds": None, "decision_makers": [],
                        "people_summary": {}, "error": str(e)[:300]}

        with concurrent.futures.ThreadPoolExecutor(max_workers=len(self.models)) as pool:
            runs = list(pool.map(one, self.models))

        ok_runs = [r for r in runs if r["status"] == 200 and r["report"]]
        if not ok_runs:
            first = runs[0] if runs else {}
            item["status"] = (STATUS_TIMEOUT if first.get("status") in (0, 504) else STATUS_FAILED)
            item["error"] = first.get("error") or "no response"
            return

        item["status"] = STATUS_PDF
        self.stage = "Generating PDF..."
        self.save()

        pdfs, total_tokens = [], 0
        for run in ok_runs:
            _record, pdf_name = write_company_output(company, website, package, run)
            pdfs.append({"model": run["model_label"], "file": pdf_name})
            total_tokens += run["total_tokens"] or 0

        item.update(status=STATUS_COMPLETE,
                    report_version=CURRENT_REPORT_VERSION,
                    limited_sources=bool(package.get("limited_evidence")),
                    model=", ".join(r["model_label"] for r in ok_runs),
                    tokens=total_tokens, pdfs=pdfs, error=None, error_reason=None,
                    report_dir=ps.safe_name(company))
        if price:
            usage = {"input": sum(r["input_tokens"] or 0 for r in ok_runs),
                     "output": sum(r["output_tokens"] or 0 for r in ok_runs)}
            item["cost"] = estimate_cost(usage, price)


def load_batch(batch_id):
    p = BATCH_DIR / ("%s.json" % batch_id)
    if not p.exists():
        return None
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except Exception:
        return None


def list_batches():
    BATCH_DIR.mkdir(parents=True, exist_ok=True)
    out = []
    for p in sorted(BATCH_DIR.glob("*.json"), key=lambda x: x.stat().st_mtime, reverse=True)[:20]:
        try:
            d = json.loads(p.read_text(encoding="utf-8"))
        except Exception:
            continue
        out.append({"id": d.get("id"), "status": d.get("status"),
                    "completed": d.get("completed"), "total": d.get("total"),
                    "saved_at": d.get("saved_at")})
    return out


def completed_pdfs():
    """Every PDF already on disk, for the ZIP download."""
    out = []
    if not REPORTS_DIR.exists():
        return out
    for meta in REPORTS_DIR.glob("*/_meta.json"):
        try:
            d = json.loads(meta.read_text(encoding="utf-8"))
        except Exception:
            continue
        for r in d.get("runs", []):
            pdf = meta.parent / r["pdf"]
            if pdf.exists():
                out.append({"company": d.get("company"), "model": r.get("model_label"),
                            "path": pdf, "file": r["pdf"],
                            "dir": meta.parent.name})
    return out


# --------------------------------------------------------------- compilation
def load_records(companies=None, order="upload"):
    """Load already-saved research records. Never calls a model.

    `companies` is an ordered list of names; `order` is upload | name | completed.
    """
    available = {}
    if REPORTS_DIR.exists():
        for meta_path in REPORTS_DIR.glob("*/_meta.json"):
            try:
                meta = json.loads(meta_path.read_text(encoding="utf-8"))
            except Exception:
                continue
            runs = meta.get("runs") or []
            if not runs:
                continue
            # Most recent run per company; one section per company.
            run = sorted(runs, key=lambda r: r.get("timestamp", ""))[-1]
            rec_path = meta_path.parent / run.get("json", "research.json")
            if not rec_path.exists():
                rec_path = meta_path.parent / "research.json"
            if not rec_path.exists():
                continue
            try:
                rec = json.loads(rec_path.read_text(encoding="utf-8"))
            except Exception:
                continue
            available[meta.get("company", meta_path.parent.name)] = rec

    if companies:
        wanted = [c for c in companies if c in available]
        records = [available[c] for c in wanted]
    else:
        records = list(available.values())

    if order == "name":
        records.sort(key=lambda r: (r.get("company") or "").lower())
    elif order == "completed":
        records.sort(key=lambda r: r.get("timestamp") or "")
    # "upload" keeps the caller-provided order (or disk order when none given)
    return records


SELECTED_TITLE = {lv.EN: "Selected Companies", lv.ZH: "所选公司研究汇总报告",
                  lv.BILINGUAL: "Selected Companies / 所选公司研究汇总报告"}


def compile_portfolio(companies=None, order="upload", cfg=None, selected=False, lang=None):
    """Build the combined PDF from saved data. Zero AI tokens.

    `lang` renders EN / ZH / bilingual from the same stored bilingual records.
    """
    lang = lv.normalize(lang)
    records = load_records(companies, order)
    if not records:
        raise ValueError("No completed reports to compile yet.")
    date = datetime.now().astimezone().strftime("%Y-%m-%d")
    stem = ("Selected_Companies_Account_Research" if selected
            else "Account_Research_Portfolio")
    out = REPORTS_DIR / ("%s_%s_%s.pdf" % (stem, lv.SUFFIX[lang], date))
    title = SELECTED_TITLE[lang] if selected else lv.label("portfolio_title", lang)
    ps.build_portfolio_pdf(records, out, title=title,
                           price=pricing(cfg) if cfg else None, lang=lang)
    return out, [r.get("company") for r in records]


# --------------------------------------------------------------- deletion
def delete_report(company, results_dir=None):
    """Delete every saved artifact for one company's report.

    The caller supplies a COMPANY NAME, never a path: the directory is derived
    server-side with ps.safe_name and then checked to be a direct child of
    reports/, so a crafted name cannot escape the store or touch another
    company's folder.

    Removed: research JSON(s), PDFs, _meta.json, history/, the cached evidence
    package (so the next run researches afresh rather than reusing old
    evidence), and matching single-run entries in test_results/.
    Kept: the company row, its name, its corrected website, the uploaded batch
    record, every other company's reports, and all credentials/configuration.
    """
    name = (company or "").strip()
    if not name:
        return {"ok": False, "error": "company is required"}

    reports_root = REPORTS_DIR.resolve()
    target = (REPORTS_DIR / ps.safe_name(name)).resolve()
    if target.parent != reports_root or target == reports_root:
        return {"ok": False, "error": "invalid company identifier"}

    removed = {"report_dir": False, "files": 0, "evidence_cache": False, "saved_results": 0}
    if target.exists() and target.is_dir():
        removed["files"] = sum(1 for _ in target.rglob("*") if _.is_file())
        shutil.rmtree(target)
        removed["report_dir"] = True

    # The evidence package is keyed per company; keeping it would let a
    # regenerated report reuse the very evidence we are trying to move past.
    try:
        for cache in rs.EVIDENCE_CACHE_DIR.glob("{}_*.json".format(
                re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")[:24] or "company")):
            cache.unlink()
            removed["evidence_cache"] = True
    except Exception:
        pass

    # Single-run history entries for the same company.
    rd = Path(results_dir) if results_dir else (HERE / "test_results")
    if rd.exists():
        for f in rd.glob("*.json"):
            try:
                if json.loads(f.read_text(encoding="utf-8")).get("company") == name:
                    f.unlink()
                    removed["saved_results"] += 1
            except Exception:
                continue

    if not removed["report_dir"] and not removed["saved_results"]:
        return {"ok": False, "error": "no saved report found for this company",
                "company": name, "removed": removed}
    return {"ok": True, "company": name, "removed": removed}


def has_report(company):
    meta = (REPORTS_DIR / ps.safe_name(company or "")) / "_meta.json"
    return meta.exists()


# --------------------------------------------------------------- language exports
def language_pdf(company, lang, model=None):
    """Path to this company's PDF in `lang`, rendering it if it does not exist yet.

    Built from the saved research JSON only - zero AI calls, zero re-retrieval.
    The result is cached on disk, so the second request is a file read.
    """
    lang = lv.normalize(lang)
    d = company_dir(company)
    prior = existing_report(company, model)
    if not prior or not prior.get("runs"):
        return None, None
    run = next((r for r in prior["runs"] if not model or r.get("model") == model),
               prior["runs"][0])
    record_path = d / (run.get("json") or "research.json")
    if not record_path.exists():
        record_path = d / "research.json"
    if not record_path.exists():
        return None, None
    try:
        record = json.loads(record_path.read_text(encoding="utf-8"))
    except Exception:
        return None, None
    when = None
    try:
        when = datetime.fromisoformat(record.get("timestamp") or "")
    except Exception:
        pass
    name = ps.pdf_filename(company, run.get("model_label"), when, multi_model=True, lang=lang)
    path = d / name
    if not path.exists():
        ps.build_pdf(record, path, lang=lang)
    return path, name


def language_pdfs(companies=None, lang=None):
    """(company, path, filename) for every completed report, in `lang`."""
    out = []
    for company in [c["company"] for c in completed_pdfs()] if companies is None else companies:
        try:
            path, name = language_pdf(company, lang)
        except Exception:
            path, name = None, None
        if path and path.exists():
            out.append({"company": company, "path": path, "file": name,
                        "dir": ps.safe_name(company)})
    return out
